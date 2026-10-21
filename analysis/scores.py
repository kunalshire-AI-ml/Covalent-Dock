"""
scores.py — read gnina scores back out of the result SDF.

score_table.py did this for one hard-coded grid of cysteines and fragments.
The reading logic is the same; the configuration is not baked in, so the same
code serves a single run in the app and a batch of runs on disk.
"""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence

AFFINITY_KEYS = ("minimizedAffinity", "Affinity", "CNNaffinity")
SCORE_COLUMNS = (
    ("pose", "Pose"),
    ("minimizedAffinity", "Affinity (kcal/mol)"),
    ("CNNscore", "CNN score"),
    ("CNNaffinity", "CNN affinity"),
    ("CNN_VS", "CNN VS"),
    ("RMSD", "RMSD"),
)


@dataclass
class Pose:
    index: int
    properties: Dict[str, str] = field(default_factory=dict)

    def number(self) -> int:
        return self.index + 1

    def get_float(self, key: str) -> Optional[float]:
        raw = self.properties.get(key)
        if raw is None:
            return None
        try:
            return float(raw.split()[0])
        except (ValueError, IndexError):
            return None

    @property
    def affinity(self) -> Optional[float]:
        for key in AFFINITY_KEYS:
            value = self.get_float(key)
            if value is not None:
                return value
        return None

    @property
    def cnn_score(self) -> Optional[float]:
        return self.get_float("CNNscore")

    @property
    def cnn_affinity(self) -> Optional[float]:
        return self.get_float("CNNaffinity")


def read_poses(sdf_path: str) -> List[Pose]:
    """Every pose in a gnina output SDF, with its score tags."""
    if not os.path.isfile(sdf_path) or os.path.getsize(sdf_path) == 0:
        return []
    poses: List[Pose] = []
    props: Dict[str, str] = {}
    key: Optional[str] = None
    index = 0
    with open(sdf_path, errors="replace") as fh:
        for line in fh:
            s = line.strip()
            if s == "$$$$":
                poses.append(Pose(index=index, properties=props))
                props, key, index = {}, None, index + 1
                continue
            if s.startswith(">") and "<" in s:
                key = s.split("<")[1].split(">")[0]
            elif key and s:
                props[key] = s
                key = None
    if props:
        poses.append(Pose(index=index, properties=props))
    return poses


def pose_table(poses: Sequence[Pose]) -> List[Dict[str, object]]:
    """Rows ready for st.dataframe / pandas."""
    rows = []
    for pose in poses:
        row: Dict[str, object] = {"Pose": pose.number()}
        for key, label in SCORE_COLUMNS[1:]:
            value = pose.get_float(key)
            if value is not None:
                row[label] = round(value, 4)
        rows.append(row)
    return rows


def best_pose(poses: Sequence[Pose], by: str = "affinity") -> Optional[Pose]:
    if not poses:
        return None
    if by == "cnn":
        scored = [p for p in poses if p.cnn_score is not None]
        return max(scored, key=lambda p: p.cnn_score) if scored else poses[0]
    scored = [p for p in poses if p.affinity is not None]
    return min(scored, key=lambda p: p.affinity) if scored else poses[0]


def summarise(poses: Sequence[Pose]) -> Dict[str, object]:
    affinities = [p.affinity for p in poses if p.affinity is not None]
    cnn = [p.cnn_score for p in poses if p.cnn_score is not None]
    return {
        "n_poses": len(poses),
        "best_affinity": min(affinities) if affinities else None,
        "mean_affinity": round(sum(affinities) / len(affinities), 3) if affinities else None,
        "best_cnn": max(cnn) if cnn else None,
    }


# ─────────────────────────────────────────────────────────────────────────────
#  Export
# ─────────────────────────────────────────────────────────────────────────────

def write_csv(rows: Sequence[Dict[str, object]], path: str,
              header_note: str = "") -> str:
    if not rows:
        raise ValueError("Nothing to export.")
    columns: List[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        if header_note:
            writer.writerow([header_note])
            writer.writerow([])
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(c, "") for c in columns])
    return path


def write_xlsx(rows: Sequence[Dict[str, object]], path: str,
               sheet_name: str = "Scores", header_note: str = "") -> str:
    """Formatted workbook; needs openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    if not rows:
        raise ValueError("Nothing to export.")
    columns: List[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    r = 1
    if header_note:
        ws.cell(row=1, column=1, value=header_note).font = Font(
            name="Arial", italic=True, size=10, color="595959")
        r = 3

    for col_i, name in enumerate(columns, start=1):
        cell = ws.cell(row=r, column=col_i, value=name)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
        ws.column_dimensions[cell.column_letter].width = max(12, len(name) + 3)

    for row_i, row in enumerate(rows, start=r + 1):
        for col_i, name in enumerate(columns, start=1):
            value = row.get(name, "")
            cell = ws.cell(row=row_i, column=col_i, value=value)
            cell.alignment = Alignment(horizontal="center")
            cell.border = box
            if isinstance(value, float):
                cell.number_format = "0.000"
            if row_i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EBF3FB")

    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
